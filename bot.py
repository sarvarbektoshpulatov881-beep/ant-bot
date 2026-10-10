"""
ANT tashrif-eslatma boti.

Rollar:
  🎧 Operator  — mijozni do'konga taklif qiladi va tashrifni kiritadi
  🛒 Sotuvchi  — mijoz keldi/kelmadi, sotib oldi/olmadi (sabab, summa) ni belgilaydi
  👑 Egasi     — hammasi + jamoani boshqarish + hisobotlar

Hammasi bot ichida (guruhsiz). Sotuvchiga (biriktirilgan yoki barcha sotuvchilarga) shaxsan:
yangi tashrif, 1 kun / 1 soat oldin eslatma, "mijoz keldimi?" eslatmasi.
Jamoaga: har kuni 09:00 da bugungi ro'yxat. Egasiga: har kuni 23:00 da kunlik hisobot.
"""
import asyncio
import csv
import html
import io
import logging
import os
import re
import secrets
import sqlite3
from contextlib import closing
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

from aiogram import Bot, Dispatcher, F, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.exceptions import TelegramMigrateToChat
from aiogram.filters import Command, CommandStart, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (BotCommand, BufferedInputFile, CallbackQuery,
                           InlineKeyboardButton, InlineKeyboardMarkup,
                           KeyboardButton, Message, ReplyKeyboardMarkup,
                           ReplyKeyboardRemove)
from dotenv import load_dotenv

load_dotenv()

# ================================================================ sozlamalar
BOT_TOKEN = os.getenv("BOT_TOKEN", "").strip()
SELLERS_CHAT_ID = int(os.getenv("SELLERS_CHAT_ID", "0").strip() or 0)
OPERATOR_IDS = {int(x) for x in os.getenv("OPERATOR_IDS", "").replace(" ", "").split(",") if x}
TZ = ZoneInfo(os.getenv("TIMEZONE", "Asia/Tashkent"))
REMINDERS = sorted(
    {int(x) for x in os.getenv("REMINDER_MINUTES", "1440,60").replace(" ", "").split(",") if x},
    reverse=True,
)
DAILY_SUMMARY_TIME = os.getenv("DAILY_SUMMARY_TIME", "09:00").strip()
REPORT_TIME = os.getenv("REPORT_TIME", "23:00").strip()           # egasiga kunlik hisobot
CHECK_AFTER_MINUTES = int(os.getenv("CHECK_AFTER_MINUTES", "60"))  # "mijoz keldimi?" eslatmasi
DB_PATH = os.getenv("DB_PATH", "ant_bot.db")
# Bot egasi. Bo'sh bo'lsa — botga birinchi /start yozgan odam egasi bo'ladi
OWNER_ID = int(os.getenv("OWNER_ID", "0").strip() or 0)

FMT = "%Y-%m-%d %H:%M"
MONTHS = ["yanvar", "fevral", "mart", "aprel", "may", "iyun", "iyul",
          "avgust", "sentabr", "oktabr", "noyabr", "dekabr"]
WEEKDAYS = ["dushanba", "seshanba", "chorshanba", "payshanba", "juma", "shanba", "yakshanba"]

REGIONS = ["Toshkent shahri", "Toshkent viloyati", "Andijon", "Buxoro", "Farg'ona", "Jizzax",
           "Xorazm", "Namangan", "Navoiy", "Qashqadaryo", "Qoraqalpog'iston", "Samarqand",
           "Sirdaryo", "Surxondaryo", "Qozog'iston", "Qirg'iziston", "Tojikiston"]
REASONS = [("narx", "💸 Narx qimmat"), ("model", "👕 Model yo'q"), ("razmer", "📏 Razmer yo'q"),
           ("keyin", "⏳ Keyinroq oladi"), ("boshqa", "❓ Boshqa sabab")]
REASON_NAMES = dict(REASONS)

# ================================================================ menyu
BTN_NEW = "🆕 Yangi mijoz"
BTN_TODAY = "📋 Bugungi tashriflar"
BTN_LIST = "🗓 Barcha tashriflar"
BTN_CANCEL = "❌ Bekor qilish"
BTN_ADMINS = "👥 Adminlar"
BTN_REPORT = "📊 Hisobot"
MENU_TEXTS = (BTN_NEW, BTN_TODAY, BTN_LIST, BTN_CANCEL, BTN_ADMINS, BTN_REPORT)


def _kb(rows):
    return ReplyKeyboardMarkup(keyboard=[[KeyboardButton(text=t) for t in r] for r in rows],
                               resize_keyboard=True, is_persistent=True)


MAIN_KB = _kb([[BTN_NEW], [BTN_TODAY, BTN_LIST], [BTN_CANCEL]])                     # operator
OWNER_KB = _kb([[BTN_NEW], [BTN_TODAY, BTN_LIST], [BTN_REPORT, BTN_ADMINS], [BTN_CANCEL]])
SELLER_KB = _kb([[BTN_TODAY, BTN_LIST]])                                             # sotuvchi

ROLE_NAMES = {"owner": "👑 Egasi", "operator": "🎧 Operator", "sotuvchi": "🛒 Sotuvchi"}

HELP = (
    "Salom! Men <b>ANT tashrif-eslatma</b> botiman. Sizning rolingiz: <b>🎧 Operator</b>\n\n"
    f"{BTN_NEW} — mijozni do'konga taklif qilib, tashrifini kiritish\n"
    f"{BTN_TODAY} — bugun keladiganlar\n"
    f"{BTN_LIST} — kelgusi tashriflar (vaqtni o'zgartirish va bekor qilish shu yerda)\n"
    f"{BTN_CANCEL} — kiritishni yarmida to'xtatish"
)
HELP_SELLER = (
    "Salom! Men <b>ANT tashrif-eslatma</b> botiman. Sizning rolingiz: <b>🛒 Sotuvchi</b>\n\n"
    "Yangi mijoz va eslatmalar sizga shu yerga keladi. Mijoz kelganda xabar ostida belgilang:\n"
    "✅ Keldi → 🛍 Sotib oldi (summani yozasiz) / 🚫 Sotib olmadi (sababini tanlaysiz)\n"
    "❌ Kelmadi\n\n"
    f"{BTN_TODAY} — bugungi mijozlar (shu yerdan ham belgilash mumkin)\n"
    f"{BTN_LIST} — kelgusi tashriflar"
)
HELP_OWNER = HELP.replace("<b>🎧 Operator</b>", "<b>👑 Egasi</b>") + (
    f"\n{BTN_REPORT} — kunlik/haftalik/oylik hisobot va Excel\n"
    f"{BTN_ADMINS} — operator va sotuvchilarni qo'shish/boshqarish\n\n"
    "Siz sotuvchi tugmalarini ham bosa olasiz. Har kuni "
    f"{REPORT_TIME} da sizga kunlik hisobot keladi."
)


# ================================================================ baza
def q(sql, params=()):
    with closing(sqlite3.connect(DB_PATH)) as con:
        con.row_factory = sqlite3.Row
        rows = con.execute(sql, params).fetchall()
        con.commit()
        return rows


def execute(sql, params=()):
    with closing(sqlite3.connect(DB_PATH)) as con:
        cur = con.execute(sql, params)
        con.commit()
        return cur.lastrowid


def init_db():
    with closing(sqlite3.connect(DB_PATH)) as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS visits(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client TEXT NOT NULL,
            phone TEXT DEFAULT '',
            visit_at TEXT NOT NULL,           -- 'YYYY-MM-DD HH:MM' (Toshkent vaqti)
            note TEXT DEFAULT '',
            status TEXT DEFAULT 'kutilmoqda', -- kutilmoqda | keldi | kelmadi | bekor
            created_by INTEGER,
            created_at TEXT,
            closed_by TEXT,
            closed_at TEXT
        );
        CREATE TABLE IF NOT EXISTS sent(visit_id INTEGER, kind TEXT, PRIMARY KEY(visit_id, kind));
        CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY, value TEXT);
        CREATE TABLE IF NOT EXISTS admins(
            user_id INTEGER PRIMARY KEY, name TEXT, username TEXT, added_at TEXT
        );
        CREATE TABLE IF NOT EXISTS requests(
            user_id INTEGER PRIMARY KEY, name TEXT, username TEXT, status TEXT, asked_at TEXT
        );
        CREATE TABLE IF NOT EXISTS invites(
            token TEXT PRIMARY KEY, role TEXT, created_at TEXT, used_by INTEGER
        );
        CREATE TABLE IF NOT EXISTS pending_amount(user_id INTEGER PRIMARY KEY, visit_id INTEGER);
        CREATE TABLE IF NOT EXISTS visit_msgs(visit_id INTEGER, chat_id INTEGER, message_id INTEGER);
        """)
        # eski bazalarga yangi ustunlar
        for table, col in (
            ("visits", "result TEXT"),            # sotib_oldi | sotib_olmadi
            ("visits", "result_by TEXT"),
            ("visits", "result_at TEXT"),
            ("visits", "result_by_id INTEGER"),
            ("visits", "reason TEXT"),            # sotib olmaslik sababi
            ("visits", "amount INTEGER"),         # savdo summasi (so'm)
            ("visits", "region TEXT"),
            ("visits", "assigned_to INTEGER"),    # biriktirilgan sotuvchi
            ("visits", "created_by_name TEXT"),
            ("visits", "phone_key TEXT"),         # takroriy mijozni topish uchun
            ("visits", "parent_id INTEGER"),      # qayta taklif qilingan tashrif
            ("admins", "role TEXT DEFAULT 'operator'"),
        ):
            try:
                con.execute(f"ALTER TABLE {table} ADD COLUMN {col}")
            except sqlite3.OperationalError:
                pass  # ustun allaqachon bor
        for vid, phone in con.execute("SELECT id, phone FROM visits WHERE phone_key IS NULL").fetchall():
            con.execute("UPDATE visits SET phone_key=? WHERE id=?", (phone_key(phone), vid))
        con.commit()


def get_visit(vid):
    rows = q("SELECT * FROM visits WHERE id=?", (vid,))
    return rows[0] if rows else None


def get_meta(key):
    rows = q("SELECT value FROM meta WHERE key=?", (key,))
    return rows[0]["value"] if rows else None


def set_meta(key, value):
    execute("INSERT OR REPLACE INTO meta(key, value) VALUES(?, ?)", (key, value))


# ================================================================ yordamchi
def g(v, k):
    """dict yoki sqlite3.Row dan xavfsiz o'qish."""
    try:
        return v[k]
    except (KeyError, IndexError):
        return None


def esc(t):
    return html.escape(str(t or ""), quote=False)


def now_local():
    return datetime.now(TZ).replace(tzinfo=None, second=0, microsecond=0)


def fmt_date(d):
    return f"{d.day}-{MONTHS[d.month - 1]} ({WEEKDAYS[d.weekday()]})"


def fmt_dt(dt):
    return f"{fmt_date(dt)}, {dt:%H:%M}"


def money(n):
    return f"{int(n or 0):,}".replace(",", " ") + " so'm"


def pct(a, b):
    return f"{round(100 * a / b)}%" if b else "—"


def phone_key(phone):
    digits = re.sub(r"\D", "", phone or "")
    return digits[-9:] if len(digits) >= 9 else ""


def human_minutes(m):
    m = max(int(m), 0)
    parts = []
    if m // 1440:
        parts.append(f"{m // 1440} kun")
    if m % 1440 // 60:
        parts.append(f"{m % 1440 // 60} soat")
    if m % 60:
        parts.append(f"{m % 60} daqiqa")
    return " ".join(parts) or "bir necha daqiqa"


def month_from_word(word):
    word = word.lower()
    for i, name in enumerate(MONTHS, 1):
        key = name[:4] if name.startswith("iyu") else name[:3]
        if word.startswith(key):
            return i
    return None


def parse_date(text, today):
    """'13.10', '13/10/2026', '13-10', '13-oktabr', '13 okt' -> date"""
    text = text.strip().lower().replace("ʻ", "'").replace("‘", "'").replace("’", "'")
    m = re.fullmatch(r"(\d{1,2})[.\-/ ](\d{1,2})(?:[.\-/ ](\d{2,4}))?", text)
    if m:
        d, mo, y = int(m[1]), int(m[2]), m[3]
    else:
        m = re.fullmatch(r"(\d{1,2})[\s\-]*([a-z']+)(?:\s+(\d{4}))?", text)
        if not m:
            return None
        d, mo, y = int(m[1]), month_from_word(m[2]), m[3]
        if not mo:
            return None
    explicit_year = y is not None
    y = int(y) if y else today.year
    if y < 100:
        y += 2000
    try:
        result = date(y, mo, d)
        if not explicit_year and result < today:
            result = date(y + 1, mo, d)
    except ValueError:
        return None
    return result


def parse_time(text):
    """'13:00', '13.00', '13;00', '13 00', '13' -> (13, 0)"""
    m = re.fullmatch(r"\s*(\d{1,2})(?:\s*[:.;,\s]\s*(\d{2}))?\s*", text)
    if not m:
        return None
    h, mi = int(m[1]), int(m[2] or 0)
    if h > 23 or mi > 59:
        return None
    return h, mi


def parse_amount(text):
    """'2500000', '2 500 000', '2,5 mln', '800 ming' -> int so'm"""
    t = text.lower().replace(" ", "").replace(",", ".")
    mult = 1_000_000 if ("mln" in t or "million" in t) else 1000 if ("ming" in t or t.endswith("k")) else 1
    if mult == 1:
        digits = re.sub(r"\D", "", text)
        val = int(digits) if digits else 0
    else:
        m = re.search(r"\d+(?:\.\d+)?", t)
        val = int(float(m.group()) * mult) if m else 0
    return val if val > 0 else None


def status_mark(v):
    st = g(v, "status")
    if st == "kelmadi":
        return " ❌ kelmadi"
    if st == "bekor":
        return " 🗑 bekor"
    if st == "keldi":
        if g(v, "result") == "sotib_oldi":
            return " 🛍 sotib oldi" + (f" — {money(v['amount'])}" if g(v, "amount") else "")
        if g(v, "result") == "sotib_olmadi":
            reason = REASON_NAMES.get(g(v, "reason"))
            return " 🚫 sotib olmadi" + (f" ({reason})" if reason else "")
        return " ✅ keldi"
    if st == "kutilmoqda" and g(v, "visit_at") and g(v, "visit_at") <= now_local().strftime(FMT):
        return " ⏳ belgilanmagan"
    return ""


# ================================================================ rollar
DENY_TEXT = "⛔ Bu bot faqat ANT jamoasi uchun."


def owner_id():
    return OWNER_ID or int(get_meta("owner_id") or 0)


def is_owner(user_id):
    return user_id == owner_id()


def admin_rows():
    return q("SELECT * FROM admins ORDER BY added_at")


def seller_rows():
    return q("SELECT * FROM admins WHERE role='sotuvchi' ORDER BY name")


def person_name(user_id):
    if not user_id:
        return ""
    rows = q("SELECT name FROM admins WHERE user_id=?", (user_id,))
    if rows:
        return rows[0]["name"]
    return "Egasi" if is_owner(user_id) else str(user_id)


def role_of(user_id):
    """owner | operator | sotuvchi | None"""
    if is_owner(user_id):
        return "owner"
    if user_id in OPERATOR_IDS:
        return "operator"
    rows = q("SELECT role FROM admins WHERE user_id=?", (user_id,))
    return (rows[0]["role"] or "operator") if rows else None


def is_admin(user_id):
    return role_of(user_id) is not None


def can_operate(user_id):
    return role_of(user_id) in ("owner", "operator")


def can_sell(user_id):
    return role_of(user_id) in ("owner", "sotuvchi")


def kb_for(user_id):
    return {"owner": OWNER_KB, "sotuvchi": SELLER_KB}.get(role_of(user_id), MAIN_KB)


def help_for(user_id):
    return {"owner": HELP_OWNER, "sotuvchi": HELP_SELLER}.get(role_of(user_id), HELP)


def user_label(name, username):
    label = esc(name or "Noma'lum")
    return f"{label} (@{username})" if username else label


async def deny(m: Message):
    await m.answer(DENY_TEXT, reply_markup=ReplyKeyboardRemove())


async def send_private(bot, user_id, text, reply_markup=None):
    if not user_id:
        return None
    try:
        return await bot.send_message(user_id, text, reply_markup=reply_markup)
    except Exception:
        logging.warning("Shaxsiy xabar yuborilmadi (botga /start yozmagan bo'lishi mumkin): %s", user_id)
        return None


async def request_access(m: Message):
    """Ruxsati yo'q odam botga yozganda — egasiga so'rov."""
    u = m.from_user
    owner = owner_id()
    if not owner:
        await deny(m)
        return
    now = now_local()
    rows = q("SELECT * FROM requests WHERE user_id=?", (u.id,))
    recent = bool(rows) and rows[0]["status"] == "kutilmoqda" and \
        now - datetime.strptime(rows[0]["asked_at"], FMT) < timedelta(minutes=10)
    if not recent:
        execute("INSERT OR REPLACE INTO requests(user_id, name, username, status, asked_at) VALUES(?,?,?,?,?)",
                (u.id, u.full_name, u.username or "", "kutilmoqda", now.strftime(FMT)))
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🎧 Operator qilish", callback_data=f"okop:{u.id}"),
             InlineKeyboardButton(text="🛒 Sotuvchi qilish", callback_data=f"oksel:{u.id}")],
            [InlineKeyboardButton(text="❌ Rad etish", callback_data=f"no:{u.id}")],
        ])
        await send_private(m.bot, owner, f"🔔 <b>Yangi so'rov</b>\n\n{user_label(u.full_name, u.username)} "
                                         f"botdan foydalanmoqchi.\nID: <code>{u.id}</code>", kb)
    await m.answer("⏳ Bu bot faqat ANT jamoasi uchun.\n"
                   "So'rovingiz bot egasiga yuborildi. Ruxsat berilgach, sizga xabar keladi.",
                   reply_markup=ReplyKeyboardRemove())


# ================================================================ jamoaga xabarlar (guruhsiz)
def team_ids(*roles):
    marks = ",".join("?" * len(roles))
    return [r["user_id"] for r in q(f"SELECT user_id FROM admins WHERE role IN ({marks})", roles)]


def visit_recipients(v):
    """Biriktirilgan sotuvchi; biriktirilmagan bo'lsa — barcha sotuvchilar; sotuvchi yo'q bo'lsa — egasi."""
    if g(v, "assigned_to"):
        return [v["assigned_to"]]
    ids = team_ids("sotuvchi")
    return ids or ([owner_id()] if owner_id() else [])


def track_msg(vid, msg):
    if msg:
        execute("INSERT INTO visit_msgs(visit_id, chat_id, message_id) VALUES(?,?,?)",
                (vid, msg.chat.id, msg.message_id))


async def clear_buttons(bot, vid, keep=None):
    """Tashrif bo'yicha yuborilgan boshqa xabarlardagi tugmalarni olib tashlash (ikki marta belgilanmasin)."""
    for r in q("SELECT * FROM visit_msgs WHERE visit_id=?", (vid,)):
        if keep is not None and r["chat_id"] == keep.chat.id and r["message_id"] == keep.message_id:
            continue
        try:
            await bot.edit_message_reply_markup(chat_id=r["chat_id"], message_id=r["message_id"], reply_markup=None)
        except Exception:
            pass
        execute("DELETE FROM visit_msgs WHERE visit_id=? AND chat_id=? AND message_id=?",
                (vid, r["chat_id"], r["message_id"]))


def visit_text(v):
    dt = datetime.strptime(g(v, "visit_at"), FMT)
    lines = [f"👤 Mijoz: <b>{esc(g(v, 'client'))}</b>"]
    if g(v, "phone"):
        lines.append(f"📞 Tel: {esc(v['phone'])}")
    if g(v, "region"):
        lines.append(f"📍 Hudud: {esc(v['region'])}")
    lines.append(f"📅 Vaqt: <b>{fmt_dt(dt)}</b>")
    if g(v, "assigned_to"):
        lines.append(f"🛒 Sotuvchi: {esc(person_name(v['assigned_to']))}")
    if g(v, "note"):
        lines.append(f"📝 Izoh: {esc(v['note'])}")
    if g(v, "created_by_name"):
        lines.append(f"🎧 Operator: {esc(v['created_by_name'])}")
    return "\n".join(lines)


def arrived_kb(vid):
    return InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ Keldi", callback_data=f"arr:{vid}"),
        InlineKeyboardButton(text="❌ Kelmadi", callback_data=f"miss:{vid}"),
    ]])


def result_kb(vid):
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🛍 Sotib oldi", callback_data=f"buy:{vid}"),
         InlineKeyboardButton(text="🚫 Sotib olmadi", callback_data=f"nobuy:{vid}")],
        [InlineKeyboardButton(text="⬅️ Orqaga", callback_data=f"back:{vid}")],
    ])


def reasons_kb(vid):
    rows = [[InlineKeyboardButton(text=name, callback_data=f"rsn:{vid}:{key}")] for key, name in REASONS]
    rows.append([InlineKeyboardButton(text="⬅️ Orqaga", callback_data=f"back2:{vid}")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


async def notify_visit(bot, vid, header, with_button=True):
    """Sotuvchi(lar)ga shaxsan: biriktirilgan bo'lsa unga, aks holda barcha sotuvchilarga."""
    v = get_visit(vid)
    text = f"{header}\n\n{visit_text(v)}"
    ok = False
    for uid in visit_recipients(v):
        msg = await send_private(bot, uid, text, arrived_kb(vid) if with_button else None)
        if msg:
            ok = True
            if with_button:
                track_msg(vid, msg)
    return ok


async def tell_creator(bot, v, text, reply_markup=None):
    if g(v, "created_by"):
        await send_private(bot, v["created_by"], text, reply_markup)


def day_summary(d):
    rows = q("SELECT * FROM visits WHERE visit_at LIKE ? AND status != 'bekor' ORDER BY visit_at",
             (d.isoformat() + "%",))
    if not rows:
        return None
    lines = [f"📋 <b>{fmt_date(d)} — tashriflar ({len(rows)} ta)</b>\n"]
    for v in rows:
        phone = f" ({esc(v['phone'])})" if v["phone"] else ""
        region = f" · {esc(v['region'])}" if g(v, "region") else ""
        lines.append(f"🕐 {v['visit_at'][11:]} — <b>{esc(v['client'])}</b>{phone}{region}{status_mark(v)}")
    return "\n".join(lines)


def client_history(key):
    """Takroriy mijoz: telefon bo'yicha avvalgi tashriflar."""
    if not key:
        return None
    rows = q("SELECT * FROM visits WHERE phone_key=? ORDER BY visit_at", (key,))
    if not rows:
        return None
    came = [v for v in rows if v["status"] == "keldi"]
    bought = [v for v in came if g(v, "result") == "sotib_oldi"]
    total = sum(g(v, "amount") or 0 for v in bought)
    last = rows[-1]
    text = (f"🔁 <b>Takroriy mijoz!</b> {esc(last['client'])}\n"
            f"Avval {len(rows)} marta yozilgan · {len(came)} marta kelgan · {len(bought)} marta sotib olgan"
            + (f" (jami {money(total)})" if total else "") +
            f"\nOxirgi tashrif: {fmt_dt(datetime.strptime(last['visit_at'], FMT))}{status_mark(last)}")
    return {"text": text, "region": g(last, "region"), "client": last["client"]}


# ================================================================ hisobotlar
MEDALS = ["🥇", "🥈", "🥉"]


def client_ranking(bought):
    """Sotib olgan mijozlar: telefon (yoki ism) bo'yicha jamlab, summa bo'yicha kamayish tartibida."""
    clients = {}
    for v in bought:
        key = g(v, "phone_key") or ("ism:" + (v["client"] or "").strip().lower())
        cl = clients.setdefault(key, {"name": v["client"], "phone": v["phone"] or "", "region": g(v, "region") or "",
                                      "deals": 0, "sum": 0, "last": v["visit_at"]})
        cl["deals"] += 1
        cl["sum"] += g(v, "amount") or 0
        if v["visit_at"] >= cl["last"]:  # eng oxirgi ma'lumot
            cl.update(name=v["client"], last=v["visit_at"], region=g(v, "region") or cl["region"],
                      phone=v["phone"] or cl["phone"])
    return sorted(clients.values(), key=lambda c: (c["sum"], c["deals"]), reverse=True)


def build_report(d1, d2):
    rows = q("SELECT * FROM visits WHERE substr(visit_at,1,10) BETWEEN ? AND ?",
             (d1.isoformat(), d2.isoformat()))
    period = fmt_date(d1) if d1 == d2 else f"{d1.day}-{MONTHS[d1.month-1]} — {d2.day}-{MONTHS[d2.month-1]}"
    title = f"📊 <b>Hisobot: {period}</b>\n"
    if not rows:
        return title + "\nBu davrda tashrif yo'q."
    now_s = now_local().strftime(FMT)
    active = [v for v in rows if v["status"] != "bekor"]
    cancelled = len(rows) - len(active)
    came = [v for v in active if v["status"] == "keldi"]
    missed = [v for v in active if v["status"] == "kelmadi"]
    unmarked = [v for v in active if v["status"] == "kutilmoqda" and v["visit_at"] <= now_s]
    upcoming = [v for v in active if v["status"] == "kutilmoqda" and v["visit_at"] > now_s]
    bought = [v for v in came if g(v, "result") == "sotib_oldi"]
    notbought = [v for v in came if g(v, "result") == "sotib_olmadi"]
    revenue = sum(g(v, "amount") or 0 for v in bought)
    with_amount = [v for v in bought if g(v, "amount")]

    out = [title,
           f"👥 Rejalashtirilgan: <b>{len(active)}</b>",
           f"✅ Keldi: <b>{len(came)}</b> ({pct(len(came), len(active))})",
           f"❌ Kelmadi: <b>{len(missed)}</b>"]
    if unmarked:
        out.append(f"⏳ Belgilanmagan: <b>{len(unmarked)}</b>")
    if upcoming:
        out.append(f"🕐 Hali kelishi kerak: <b>{len(upcoming)}</b>")
    if cancelled:
        out.append(f"🗑 Bekor qilingan: {cancelled}")
    out += ["",
            f"🛍 Sotib oldi: <b>{len(bought)}</b> · 🚫 Sotib olmadi: <b>{len(notbought)}</b>",
            f"📈 Konversiya (kelganlardan sotib olgan): <b>{pct(len(bought), len(came))}</b>",
            f"💰 Tushum: <b>{money(revenue)}</b>"]
    if with_amount:
        out.append(f"🧾 O'rtacha savdo: {money(revenue // len(with_amount))}")

    # mijozlar reytingi — eng ko'p xarid qilganlar yuqorida
    top = client_ranking(bought)
    if top:
        out += ["", "🏆 <b>Eng ko'p xarid qilgan mijozlar</b>"]
        for i, cl in enumerate(top[:15]):
            medal = MEDALS[i] if i < 3 else f"{i + 1}."
            extra = " · ".join(x for x in (esc(cl["region"]), esc(cl["phone"])) if x)
            amount = money(cl["sum"]) if cl["sum"] else "summa kiritilmagan"
            out.append(f"{medal} <b>{esc(cl['name'])}</b> — {amount} ({cl['deals']} ta xarid)"
                       + (f"\n      {extra}" if extra else ""))
        if len(top) > 15:
            out.append(f"... va yana {len(top) - 15} ta mijoz (to'liq ro'yxat Excelda)")

    # operatorlar
    ops = {}
    for v in active:
        oid = g(v, "created_by")
        o = ops.setdefault(oid, {"name": g(v, "created_by_name") or person_name(oid) or "?",
                                 "planned": 0, "came": 0, "bought": 0, "sum": 0})
        o["planned"] += 1
        if v["status"] == "keldi":
            o["came"] += 1
            if g(v, "result") == "sotib_oldi":
                o["bought"] += 1
                o["sum"] += g(v, "amount") or 0
    out += ["", "🎧 <b>Operatorlar</b>"]
    ranked = sorted(ops.values(), key=lambda o: (o["sum"], o["bought"], o["came"], o["planned"]), reverse=True)
    for i, o in enumerate(ranked):
        medal = MEDALS[i] if i < 3 else f"{i + 1}."
        out.append(f"{medal} {esc(o['name'])} — {o['planned']} taklif · {o['came']} keldi · "
                   f"{o['bought']} sotib oldi · {money(o['sum'])}")

    # hududlar
    regions = {}
    for v in active:
        r = regions.setdefault(g(v, "region") or "Ko'rsatilmagan", [0, 0, 0])
        r[0] += 1
        if v["status"] == "keldi":
            r[1] += 1
            if g(v, "result") == "sotib_oldi":
                r[2] += 1
    out += ["", "📍 <b>Hududlar</b>"]
    for name, (p, c, b) in sorted(regions.items(), key=lambda x: (x[1][2], x[1][1], x[1][0]), reverse=True):
        out.append(f"• {esc(name)} — {p} taklif · {c} keldi · {b} sotib oldi")

    # sabablar
    if notbought:
        out += ["", "🚫 <b>Sotib olmaslik sabablari</b>"]
        counts = {}
        for v in notbought:
            counts[g(v, "reason") or "boshqa"] = counts.get(g(v, "reason") or "boshqa", 0) + 1
        for key, n in sorted(counts.items(), key=lambda x: x[1], reverse=True):
            out.append(f"{REASON_NAMES.get(key, key)} — {n}")
    return "\n".join(out)


STATUS_UZ = {"kutilmoqda": "Kutilmoqda", "keldi": "Keldi", "kelmadi": "Kelmadi", "bekor": "Bekor qilingan"}
RESULT_UZ = {"sotib_oldi": "Sotib oldi", "sotib_olmadi": "Sotib olmadi"}


def build_excel():
    rows = q("SELECT * FROM visits ORDER BY visit_at DESC")
    headers = ["№", "Mijoz", "Telefon", "Hudud", "Sana", "Vaqt", "Holat", "Natija", "Sabab",
               "Summa (so'm)", "Operator", "Biriktirilgan sotuvchi", "Belgilagan sotuvchi", "Izoh", "Kiritilgan"]
    data = []
    for v in rows:
        data.append([
            v["id"], v["client"], v["phone"] or "", g(v, "region") or "", v["visit_at"][:10], v["visit_at"][11:],
            STATUS_UZ.get(v["status"], v["status"]), RESULT_UZ.get(g(v, "result"), ""),
            REASON_NAMES.get(g(v, "reason"), "").split(" ", 1)[-1] if g(v, "reason") else "",
            g(v, "amount"), g(v, "created_by_name") or person_name(g(v, "created_by")),
            person_name(g(v, "assigned_to")), g(v, "result_by") or "", v["note"] or "", v["created_at"] or "",
        ])
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Tashriflar"
        ws.append(headers)
        for r in data:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F5597")
        widths = [6, 28, 18, 18, 12, 8, 14, 14, 18, 16, 20, 20, 20, 30, 17]
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        for c in ws["J"][1:]:
            c.number_format = "#,##0"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        ws2 = wb.create_sheet("Mijozlar reytingi")
        ws2.append(["O'rin", "Mijoz", "Telefon", "Hudud", "Xaridlar soni", "Jami summa (so'm)", "Oxirgi xarid"])
        bought = [v for v in rows if v["status"] == "keldi" and g(v, "result") == "sotib_oldi"]
        for i, cl in enumerate(client_ranking(bought), 1):
            ws2.append([i, cl["name"], cl["phone"], cl["region"], cl["deals"], cl["sum"], cl["last"][:10]])
        for cell in ws2[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F5597")
        for col, w in zip("ABCDEFG", (7, 28, 18, 18, 14, 20, 14)):
            ws2.column_dimensions[col].width = w
        for c in ws2["F"][1:]:
            c.number_format = "#,##0"
        ws2.freeze_panes = "A2"
        bio = io.BytesIO()
        wb.save(bio)
        return bio.getvalue(), "xlsx"
    except ImportError:
        sio = io.StringIO()
        w = csv.writer(sio, delimiter=";")
        w.writerow(headers)
        w.writerows(data)
        return sio.getvalue().encode("utf-8-sig"), "csv"


def report_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="Bugun", callback_data="rep:today"),
         InlineKeyboardButton(text="Kecha", callback_data="rep:yday")],
        [InlineKeyboardButton(text="Oxirgi 7 kun", callback_data="rep:week"),
         InlineKeyboardButton(text="Shu oy", callback_data="rep:month")],
        [InlineKeyboardButton(text="📥 Excel (tashriflar + mijozlar reytingi)", callback_data="xls")],
    ])


# ================================================================ holatlar
class NewVisit(StatesGroup):
    client = State()
    phone = State()
    region = State()
    region_other = State()
    vdate = State()
    vtime = State()
    note = State()
    seller = State()
    confirm = State()


class Resched(StatesGroup):
    vdate = State()
    vtime = State()


router = Router()
# kiritish bosqichlarida oddiy matn (menyu tugmasi yoki /buyruq emas)
TEXT = F.text & ~F.text.in_(MENU_TEXTS) & ~F.text.startswith("/")


def date_kb():
    t = now_local().date()
    rows = []
    for label, d in (("Bugun", t), ("Ertaga", t + timedelta(1)), ("Indinga", t + timedelta(2))):
        rows.append([InlineKeyboardButton(text=f"{label}, {d.day}-{MONTHS[d.month - 1]}",
                                          callback_data=f"d:{d.isoformat()}")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def region_kb(prev=None):
    rows = []
    if prev:
        rows.append([InlineKeyboardButton(text=f"📍 {prev} (avvalgi)", callback_data="reg:prev")])
    for i in range(0, len(REGIONS), 2):
        rows.append([InlineKeyboardButton(text=REGIONS[j], callback_data=f"reg:{j}")
                     for j in range(i, min(i + 2, len(REGIONS)))])
    rows.append([InlineKeyboardButton(text="✏️ Boshqa (yozish)", callback_data="reg:other")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


# ================================================================ umumiy buyruqlar
@router.message(Command("id"))
async def cmd_id(m: Message):
    await m.answer(f"Chat ID: <code>{m.chat.id}</code>\nSizning ID: <code>{m.from_user.id}</code>")


async def use_invite(m: Message, token):
    rows = q("SELECT * FROM invites WHERE token=?", (token,))
    if not rows or rows[0]["used_by"] or \
            now_local() - datetime.strptime(rows[0]["created_at"], FMT) > timedelta(hours=24):
        await m.answer("⚠️ Bu taklif havolasi eskirgan yoki allaqachon ishlatilgan. "
                       "Bot egasidan yangi havola so'rang.")
        return False
    u = m.from_user
    role = rows[0]["role"]
    execute("INSERT OR REPLACE INTO admins(user_id, name, username, added_at, role) VALUES(?,?,?,?,?)",
            (u.id, u.full_name, u.username or "", now_local().strftime(FMT), role))
    execute("UPDATE invites SET used_by=? WHERE token=?", (u.id, token))
    await send_private(m.bot, owner_id(), f"✅ {user_label(u.full_name, u.username)} taklif havolasi orqali "
                                          f"{ROLE_NAMES[role]} sifatida qo'shildi.")
    return True


@router.message(CommandStart(), F.chat.type == "private")
async def cmd_start(m: Message, state: FSMContext):
    await state.clear()
    uid = m.from_user.id
    if not owner_id():
        set_meta("owner_id", str(uid))
        logging.info("Bot egasi belgilandi: %s", uid)
        await m.answer("👑 Siz botning <b>egasi</b> sifatida belgilandingiz.\n"
                       f"Jamoani boshqarish «{BTN_ADMINS}» bo'limida.")
    parts = (m.text or "").split(maxsplit=1)
    if len(parts) > 1 and parts[1].startswith("inv_") and not is_owner(uid):
        await use_invite(m, parts[1][4:])
    if not is_admin(uid):
        await request_access(m)
        return
    await m.answer(help_for(uid), reply_markup=kb_for(uid))


@router.message(Command("bekor"), F.chat.type == "private")
@router.message(F.text == BTN_CANCEL, F.chat.type == "private")
async def cmd_cancel(m: Message, state: FSMContext):
    await state.clear()
    await m.answer("Kiritish to'xtatildi.", reply_markup=kb_for(m.from_user.id))


# ================================================================ operator: yangi tashrif
@router.message(Command("yangi"), F.chat.type == "private")
@router.message(F.text == BTN_NEW, F.chat.type == "private")
async def cmd_new(m: Message, state: FSMContext):
    if not can_operate(m.from_user.id):
        if is_admin(m.from_user.id):
            await m.answer("Mijoz kiritish — operator vazifasi. Sizning vazifangiz: mijoz kelganda "
                           "Keldi/Kelmadi va Sotib oldi/olmadi ni belgilash.", reply_markup=kb_for(m.from_user.id))
        else:
            await deny(m)
        return
    await state.clear()
    await state.set_state(NewVisit.client)
    await m.answer("1/6 · Mijozning <b>ism familiyasini</b> yozing (yoki do'kon nomini):")


@router.message(NewVisit.client, TEXT)
async def st_client(m: Message, state: FSMContext):
    await state.update_data(client=m.text.strip()[:200])
    await state.set_state(NewVisit.phone)
    await m.answer("2/6 · <b>Telefon raqami</b> (masalan: +998 90 123 45 67).\nRaqam bo'lmasa <b>-</b> yozing:")


@router.message(NewVisit.phone, TEXT)
async def st_phone(m: Message, state: FSMContext):
    phone = m.text.strip()
    if phone == "-":
        phone, key = "", ""
    else:
        key = phone_key(phone)
        if not key:
            await m.answer("Raqam to'liq emas. Kamida 9 ta raqam bo'lsin, masalan: <b>+998 90 123 45 67</b>")
            return
    hist = client_history(key)
    await state.update_data(phone=phone[:50], phone_key=key, prev_region=hist["region"] if hist else None)
    if hist:
        await m.answer(hist["text"])
    await state.set_state(NewVisit.region)
    await m.answer("3/6 · Mijoz <b>qayerdan</b>?", reply_markup=region_kb(hist["region"] if hist else None))


async def ask_date(msg: Message, state: FSMContext, step="4/6 · "):
    await msg.answer(f"{step}<b>Qaysi kuni</b> keladi? Tugmani bosing yoki yozing, masalan: "
                     "<b>13.10</b> yoki <b>13-oktabr</b>", reply_markup=date_kb())


@router.callback_query(NewVisit.region, F.data.startswith("reg:"))
async def st_region_btn(c: CallbackQuery, state: FSMContext):
    key = c.data[4:]
    await c.answer()
    await c.message.edit_reply_markup(reply_markup=None)
    if key == "other":
        await state.set_state(NewVisit.region_other)
        await c.message.answer("Hududni yozing (masalan: Termiz, Olmaota):")
        return
    region = (await state.get_data()).get("prev_region") if key == "prev" else REGIONS[int(key)]
    await state.update_data(region=region)
    await c.message.answer(f"📍 {esc(region)}")
    await state.set_state(NewVisit.vdate)
    await ask_date(c.message, state)


@router.message(NewVisit.region_other, TEXT)
async def st_region_other(m: Message, state: FSMContext):
    await state.update_data(region=m.text.strip()[:60])
    await state.set_state(NewVisit.vdate)
    await ask_date(m, state)


async def on_date(msg: Message, state: FSMContext, d: date):
    await state.update_data(vdate=d.isoformat())
    cur = await state.get_state()
    nxt = NewVisit.vtime if cur == NewVisit.vdate.state else Resched.vtime
    await state.set_state(nxt)
    step = "5/6 · " if nxt is NewVisit.vtime else ""
    await msg.answer(f"📅 {fmt_date(d)}\n{step}<b>Soat nechida</b> keladi? Masalan: <b>13:00</b>")


@router.callback_query(StateFilter(NewVisit.vdate, Resched.vdate), F.data.startswith("d:"))
async def st_date_btn(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await c.message.edit_reply_markup(reply_markup=None)
    await on_date(c.message, state, date.fromisoformat(c.data[2:]))


@router.message(StateFilter(NewVisit.vdate, Resched.vdate), TEXT)
async def st_date(m: Message, state: FSMContext):
    today = now_local().date()
    d = parse_date(m.text, today)
    if not d:
        await m.answer("Sanani tushunmadim. Masalan: <b>13.10</b> yoki <b>13-oktabr</b>")
        return
    if d < today:
        await m.answer("Bu sana o'tib ketgan. Boshqa sana yozing:")
        return
    await on_date(m, state, d)


async def read_time(m: Message, state: FSMContext):
    t = parse_time(m.text)
    if not t:
        await m.answer("Vaqtni tushunmadim. Masalan: <b>13:00</b>")
        return None
    d = date.fromisoformat((await state.get_data())["vdate"])
    dt = datetime(d.year, d.month, d.day, *t)
    if dt <= now_local():
        await m.answer("Bu vaqt o'tib ketgan. Boshqa vaqt yozing:")
        return None
    return dt


@router.message(NewVisit.vtime, TEXT)
async def st_time(m: Message, state: FSMContext):
    dt = await read_time(m, state)
    if not dt:
        return
    await state.update_data(visit_at=dt.strftime(FMT))
    await state.set_state(NewVisit.note)
    await m.answer("6/6 · <b>Izoh</b> (masalan: qaysi tovarlarni ko'rmoqchi). Bo'lmasa <b>-</b> yozing:")


def confirm_kb():
    return InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ Saqlash", callback_data="save"),
        InlineKeyboardButton(text="❌ Bekor", callback_data="drop"),
    ]])


async def show_confirm(msg: Message, state: FSMContext, user):
    await state.update_data(created_by_name=user.full_name)
    await state.set_state(NewVisit.confirm)
    await msg.answer("Tekshiring:\n\n" + visit_text(await state.get_data()), reply_markup=confirm_kb())


@router.message(NewVisit.note, TEXT)
async def st_note(m: Message, state: FSMContext):
    note = m.text.strip()
    await state.update_data(note="" if note == "-" else note[:500], assigned_to=None)
    sellers = seller_rows()
    if not sellers:
        await show_confirm(m, state, m.from_user)
        return
    rows = [[InlineKeyboardButton(text=f"🛒 {s['name']}", callback_data=f"sel:{s['user_id']}")] for s in sellers]
    rows.append([InlineKeyboardButton(text="👥 Biriktirmaslik (hamma sotuvchilar)", callback_data="sel:0")])
    await state.set_state(NewVisit.seller)
    await m.answer("Mijozni <b>qaysi sotuvchiga</b> biriktiramiz?",
                   reply_markup=InlineKeyboardMarkup(inline_keyboard=rows))


@router.callback_query(NewVisit.seller, F.data.startswith("sel:"))
async def st_seller(c: CallbackQuery, state: FSMContext):
    sid = int(c.data[4:])
    await c.answer()
    await c.message.edit_reply_markup(reply_markup=None)
    await state.update_data(assigned_to=sid or None)
    await show_confirm(c.message, state, c.from_user)


@router.callback_query(NewVisit.confirm, F.data == "save")
async def st_save(c: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    await state.clear()
    vid = execute(
        "INSERT INTO visits(client, phone, phone_key, region, visit_at, note, assigned_to, "
        "created_by, created_by_name, created_at) VALUES(?,?,?,?,?,?,?,?,?,?)",
        (data["client"], data.get("phone", ""), data.get("phone_key", ""), data.get("region"),
         data["visit_at"], data.get("note", ""), data.get("assigned_to"),
         c.from_user.id, c.from_user.full_name, now_local().strftime(FMT)),
    )
    await c.answer("Saqlandi")
    await c.message.edit_reply_markup(reply_markup=None)
    ok = await notify_visit(c.bot, vid, "🆕 <b>Yangi tashrif belgilandi</b>")
    await c.message.answer(
        f"✅ Saqlandi (№{vid}). " + ("Sotuvchiga xabar yuborildi." if ok
                                     else "⚠️ Sotuvchiga yuborilmadi: jamoada sotuvchi yo'q yoki u botga /start yozmagan."),
        reply_markup=kb_for(c.from_user.id))


@router.callback_query(NewVisit.confirm, F.data == "drop")
async def st_drop(c: CallbackQuery, state: FSMContext):
    await state.clear()
    await c.answer()
    await c.message.edit_reply_markup(reply_markup=None)
    await c.message.answer(f"Bekor qilindi. Yangi tashrif uchun «{BTN_NEW}» ni bosing.",
                           reply_markup=kb_for(c.from_user.id))


# ================================================================ vaqtni o'zgartirish / qayta taklif
@router.callback_query(F.data.startswith("mv:"))
async def cb_move(c: CallbackQuery, state: FSMContext):
    if not can_operate(c.from_user.id):
        await c.answer("Faqat operator o'zgartira oladi", show_alert=True)
        return
    vid = int(c.data[3:])
    v = get_visit(vid)
    if not v or v["status"] != "kutilmoqda":
        await c.answer("Bu tashrifni o'zgartirib bo'lmaydi", show_alert=True)
        return
    await c.answer()
    await state.clear()
    await state.set_state(Resched.vdate)
    await state.update_data(mode="move", vid=vid)
    await c.message.answer(f"🕐 №{vid} — <b>{esc(v['client'])}</b>\nYangi sanani tanlang:", reply_markup=date_kb())


@router.callback_query(F.data.startswith("fu:"))
async def cb_followup(c: CallbackQuery, state: FSMContext):
    _, action, vid = c.data.split(":")
    vid = int(vid)
    v = get_visit(vid)
    await c.answer()
    await c.message.edit_reply_markup(reply_markup=None)
    if action == "no" or not v:
        await c.message.answer("Yaxshi, qayta taklif qilinmaydi.")
        return
    if not can_operate(c.from_user.id):
        return
    await state.clear()
    await state.set_state(Resched.vdate)
    await state.update_data(mode="again", vid=vid)
    await c.message.answer(f"🔁 <b>{esc(v['client'])}</b> uchun yangi tashrif sanasini tanlang:",
                           reply_markup=date_kb())


@router.message(Resched.vtime, TEXT)
async def st_resched_time(m: Message, state: FSMContext):
    dt = await read_time(m, state)
    if not dt:
        return
    data = await state.get_data()
    await state.clear()
    v = get_visit(data["vid"])
    new_at = dt.strftime(FMT)
    if data["mode"] == "move":
        execute("UPDATE visits SET visit_at=? WHERE id=?", (new_at, v["id"]))
        execute("DELETE FROM sent WHERE visit_id=?", (v["id"],))
        await m.answer(f"✅ №{v['id']} vaqti o'zgartirildi: {fmt_dt(dt)}", reply_markup=kb_for(m.from_user.id))
        await notify_visit(m.bot, v["id"], f"🕐 <b>Tashrif vaqti o'zgartirildi</b> (avval: "
                                           f"{fmt_dt(datetime.strptime(v['visit_at'], FMT))})")
    else:
        vid = execute(
            "INSERT INTO visits(client, phone, phone_key, region, visit_at, note, assigned_to, "
            "created_by, created_by_name, created_at, parent_id) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
            (v["client"], v["phone"], g(v, "phone_key"), g(v, "region"), new_at, v["note"],
             g(v, "assigned_to"), m.from_user.id, m.from_user.full_name, now_local().strftime(FMT), v["id"]))
        await m.answer(f"✅ Yangi tashrif saqlandi (№{vid}): {fmt_dt(dt)}", reply_markup=kb_for(m.from_user.id))
        await notify_visit(m.bot, vid, "🔁 <b>Qayta taklif: yangi tashrif belgilandi</b>")


# ================================================================ ro'yxatlar
@router.message(Command("bugun"))
@router.message(F.text == BTN_TODAY, F.chat.type == "private")
async def cmd_today(m: Message):
    private = m.chat.type == "private"
    if private and not is_admin(m.from_user.id):
        await deny(m)
        return
    today = now_local().date()
    await m.answer(day_summary(today) or "Bugun tashrif rejalashtirilmagan.")
    if private and can_sell(m.from_user.id):
        rows = q("SELECT * FROM visits WHERE visit_at LIKE ? AND "
                 "(status='kutilmoqda' OR (status='keldi' AND result IS NULL)) ORDER BY visit_at",
                 (today.isoformat() + "%",))
        for v in rows:
            kb = arrived_kb(v["id"]) if v["status"] == "kutilmoqda" else result_kb(v["id"])
            track_msg(v["id"], await m.answer("🛒 <b>Belgilash kerak</b>\n\n" + visit_text(v), reply_markup=kb))


@router.message(Command("royxat"))
@router.message(F.text == BTN_LIST, F.chat.type == "private")
async def cmd_list(m: Message):
    if m.chat.type == "private" and not is_admin(m.from_user.id):
        await deny(m)
        return
    since = (now_local() - timedelta(hours=3)).strftime(FMT)
    rows = q("SELECT * FROM visits WHERE status='kutilmoqda' AND visit_at >= ? ORDER BY visit_at LIMIT 30",
             (since,))
    if not rows:
        await m.answer("Kelgusi tashriflar yo'q.")
        return
    lines = ["📋 <b>Kelgusi tashriflar</b>\n"]
    for v in rows:
        dt = datetime.strptime(v["visit_at"], FMT)
        region = f" · {esc(v['region'])}" if g(v, "region") else ""
        lines.append(f"№{v['id']} · {fmt_dt(dt)} — <b>{esc(v['client'])}</b>{region}")
    kb = None
    if m.chat.type == "private" and can_operate(m.from_user.id):
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text=f"🕐 №{v['id']} vaqtini o'zgartirish", callback_data=f"mv:{v['id']}"),
             InlineKeyboardButton(text=f"❌ №{v['id']} bekor", callback_data=f"cancel:{v['id']}")]
            for v in rows
        ])
    await m.answer("\n".join(lines), reply_markup=kb)


@router.callback_query(F.data.startswith("cancel:"))
async def cb_cancel(c: CallbackQuery):
    if not can_operate(c.from_user.id):
        await c.answer("Tashrifni faqat operator bekor qila oladi", show_alert=True)
        return
    vid = int(c.data.split(":")[1])
    v = get_visit(vid)
    if not v or v["status"] != "kutilmoqda":
        await c.answer("Bu tashrif allaqachon yopilgan", show_alert=True)
        return
    execute("UPDATE visits SET status='bekor', closed_by=?, closed_at=? WHERE id=?",
            (c.from_user.full_name, now_local().strftime(FMT), vid))
    await c.answer("Bekor qilindi")
    await c.message.answer(f"❌ №{vid} bekor qilindi.")
    await clear_buttons(c.bot, vid)
    await notify_visit(c.bot, vid, "❌ <b>Tashrif bekor qilindi</b>", with_button=False)


# ================================================================ sotuvchi: keldi / sotib oldi
SELL_ONLY = "Bu tugmani faqat sotuvchilar bosa oladi"


@router.callback_query(F.data.startswith("arr:") | F.data.startswith("miss:"))
async def cb_came(c: CallbackQuery):
    if not can_sell(c.from_user.id):
        await c.answer(SELL_ONLY, show_alert=True)
        return
    action, vid = c.data.split(":")
    vid = int(vid)
    v = get_visit(vid)
    if not v:
        await c.answer("Tashrif topilmadi", show_alert=True)
        return
    if v["status"] != "kutilmoqda":
        await c.answer(f"Allaqachon belgilangan: {status_mark(v).strip() or v['status']}", show_alert=True)
        return
    new_status = "keldi" if action == "arr" else "kelmadi"
    execute("UPDATE visits SET status=?, closed_by=?, closed_at=? WHERE id=?",
            (new_status, c.from_user.full_name, now_local().strftime(FMT), vid))
    await c.answer("Belgilandi")
    await clear_buttons(c.bot, vid, keep=c.message)
    if new_status == "keldi":
        await c.message.edit_reply_markup(reply_markup=result_kb(vid))
        return
    await c.message.edit_reply_markup(reply_markup=None)
    client = esc(v["client"])
    await c.message.reply(f"❌ Belgilandi: mijoz <b>{client}</b> kelmadi.")
    fu_kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="📅 Qayta taklif qilish", callback_data=f"fu:new:{vid}"),
        InlineKeyboardButton(text="✖️ Kerak emas", callback_data=f"fu:no:{vid}"),
    ]])
    await tell_creator(c.bot, v, f"❌ №{vid} — <b>{client}</b> kelmadi.\n"
                                 f"📞 {esc(v['phone']) or 'raqam yoq'}\n\n"
                                 "Mijozga qayta qo'ng'iroq qilib, yangi vaqt belgilaysizmi?", fu_kb)


@router.callback_query(F.data.startswith("buy:") | F.data.startswith("nobuy:"))
async def cb_result(c: CallbackQuery):
    if not can_sell(c.from_user.id):
        await c.answer(SELL_ONLY, show_alert=True)
        return
    action, vid = c.data.split(":")
    vid = int(vid)
    v = get_visit(vid)
    if not v:
        await c.answer("Tashrif topilmadi", show_alert=True)
        return
    if v["status"] != "keldi":
        await c.answer("Avval «✅ Keldi» belgilanishi kerak", show_alert=True)
        return
    if g(v, "result"):
        await c.answer(f"Allaqachon belgilangan:{status_mark(v)}", show_alert=True)
        return
    if action == "nobuy":
        await c.answer()
        await c.message.edit_reply_markup(reply_markup=reasons_kb(vid))
        return
    execute("UPDATE visits SET result='sotib_oldi', result_by=?, result_by_id=?, result_at=? WHERE id=?",
            (c.from_user.full_name, c.from_user.id, now_local().strftime(FMT), vid))
    await c.answer("Belgilandi")
    await c.message.edit_reply_markup(reply_markup=None)
    client = esc(v["client"])
    await clear_buttons(c.bot, vid)
    await c.message.reply(f"🛍 Belgilandi: mijoz <b>{client}</b> keldi va mahsulot sotib oldi.")
    await tell_creator(c.bot, v, f"🛍 №{vid} — <b>{client}</b> keldi va mahsulot sotib oldi.")
    # sotuvchidan savdo summasini so'raymiz (shaxsiy chatda)
    execute("INSERT OR REPLACE INTO pending_amount(user_id, visit_id) VALUES(?, ?)", (c.from_user.id, vid))
    skip = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="⏭ O'tkazib yuborish", callback_data=f"skipamt:{vid}")]])
    await send_private(c.bot, c.from_user.id,
                       f"💰 <b>{client}</b> — savdo summasini yozing (so'mda).\n"
                       "Masalan: <b>2500000</b>, <b>2,5 mln</b> yoki <b>800 ming</b>", skip)


@router.callback_query(F.data.startswith("rsn:"))
async def cb_reason(c: CallbackQuery):
    if not can_sell(c.from_user.id):
        await c.answer(SELL_ONLY, show_alert=True)
        return
    _, vid, key = c.data.split(":")
    vid = int(vid)
    v = get_visit(vid)
    if not v or v["status"] != "keldi" or g(v, "result"):
        await c.answer("Allaqachon belgilangan", show_alert=True)
        return
    execute("UPDATE visits SET result='sotib_olmadi', reason=?, result_by=?, result_by_id=?, result_at=? "
            "WHERE id=?", (key, c.from_user.full_name, c.from_user.id, now_local().strftime(FMT), vid))
    await c.answer("Belgilandi")
    await c.message.edit_reply_markup(reply_markup=None)
    client = esc(v["client"])
    text = f"mijoz <b>{client}</b> keldi, lekin sotib olmadi.\nSabab: {REASON_NAMES.get(key, key)}"
    await clear_buttons(c.bot, vid)
    await c.message.reply(f"🚫 Belgilandi: {text}")
    await tell_creator(c.bot, v, f"🚫 №{vid} — {text}")


@router.callback_query(F.data.startswith("back2:"))
async def cb_back_reasons(c: CallbackQuery):
    if not can_sell(c.from_user.id):
        await c.answer(SELL_ONLY, show_alert=True)
        return
    vid = int(c.data.split(":")[1])
    await c.answer()
    await c.message.edit_reply_markup(reply_markup=result_kb(vid))


@router.callback_query(F.data.startswith("back:"))
async def cb_back(c: CallbackQuery):
    """«Keldi» xato bosilgan bo'lsa — 1-bosqichga qaytish."""
    if not can_sell(c.from_user.id):
        await c.answer(SELL_ONLY, show_alert=True)
        return
    vid = int(c.data.split(":")[1])
    v = get_visit(vid)
    if not v or v["status"] != "keldi" or g(v, "result"):
        await c.answer("Qaytarib bo'lmaydi", show_alert=True)
        return
    execute("UPDATE visits SET status='kutilmoqda', closed_by=NULL, closed_at=NULL WHERE id=?", (vid,))
    await c.answer()
    await c.message.edit_reply_markup(reply_markup=arrived_kb(vid))


@router.callback_query(F.data.startswith("skipamt:"))
async def cb_skip_amount(c: CallbackQuery):
    execute("DELETE FROM pending_amount WHERE user_id=?", (c.from_user.id,))
    await c.answer("O'tkazib yuborildi")
    await c.message.edit_reply_markup(reply_markup=None)


# ================================================================ egasi: hisobot
@router.message(Command("hisobot"), F.chat.type == "private")
@router.message(F.text == BTN_REPORT, F.chat.type == "private")
async def cmd_report(m: Message):
    if not is_owner(m.from_user.id):
        await m.answer("Hisobotlar faqat bot egasi uchun.") if is_admin(m.from_user.id) else await deny(m)
        return
    d = now_local().date()
    await m.answer(build_report(d, d), reply_markup=report_kb())


@router.callback_query(F.data.startswith("rep:"))
async def cb_report(c: CallbackQuery):
    if not is_owner(c.from_user.id):
        await c.answer("Faqat bot egasi uchun", show_alert=True)
        return
    t = now_local().date()
    period = c.data[4:]
    d1, d2 = {"today": (t, t), "yday": (t - timedelta(1), t - timedelta(1)),
              "week": (t - timedelta(6), t), "month": (t.replace(day=1), t)}.get(period, (t, t))
    await c.answer()
    await c.message.answer(build_report(d1, d2), reply_markup=report_kb())


@router.callback_query(F.data == "xls")
async def cb_excel(c: CallbackQuery):
    if not is_owner(c.from_user.id):
        await c.answer("Faqat bot egasi uchun", show_alert=True)
        return
    await c.answer("Tayyorlanmoqda...")
    data, ext = build_excel()
    name = f"ANT_tashriflar_{now_local():%Y-%m-%d}.{ext}"
    await c.message.answer_document(BufferedInputFile(data, filename=name),
                                    caption="📥 1-varaq: barcha tashriflar (filtr bilan). 2-varaq: eng ko'p xarid qilgan mijozlar reytingi.")


# ================================================================ egasi: jamoa
@router.callback_query(F.data.startswith("ok:") | F.data.startswith("okop:")
                       | F.data.startswith("oksel:") | F.data.startswith("no:"))
async def cb_request(c: CallbackQuery):
    if not is_owner(c.from_user.id):
        await c.answer("Faqat bot egasi hal qila oladi", show_alert=True)
        return
    action, uid = c.data.split(":")
    uid = int(uid)
    rows = q("SELECT * FROM requests WHERE user_id=?", (uid,))
    name, username = (rows[0]["name"], rows[0]["username"]) if rows else ("Noma'lum", "")
    label = user_label(name, username)
    await c.answer()
    await c.message.edit_reply_markup(reply_markup=None)
    if action in ("ok", "okop", "oksel"):
        role = "sotuvchi" if action == "oksel" else "operator"
        execute("INSERT OR REPLACE INTO admins(user_id, name, username, added_at, role) VALUES(?,?,?,?,?)",
                (uid, name, username, now_local().strftime(FMT), role))
        execute("UPDATE requests SET status='ruxsat' WHERE user_id=?", (uid,))
        await c.message.answer(f"✅ {label} — {ROLE_NAMES[role]} sifatida qo'shildi.")
        await send_private(c.bot, uid, "✅ Sizga botdan foydalanishga ruxsat berildi!\n\n" + help_for(uid),
                           kb_for(uid))
    else:
        execute("UPDATE requests SET status='rad' WHERE user_id=?", (uid,))
        await c.message.answer(f"❌ {label} ning so'rovi rad etildi.")
        await send_private(c.bot, uid, "❌ Botdan foydalanish so'rovingiz rad etildi.")


def team_kb(rows):
    buttons = [[InlineKeyboardButton(text="➕ Operator qo'shish", callback_data="inv:operator"),
                InlineKeyboardButton(text="➕ Sotuvchi qo'shish", callback_data="inv:sotuvchi")]]
    for r in rows:
        other = "sotuvchi" if (r["role"] or "operator") == "operator" else "operator"
        buttons.append([
            InlineKeyboardButton(text=f"🔄 {r['name']} → {ROLE_NAMES[other]}", callback_data=f"role:{r['user_id']}"),
            InlineKeyboardButton(text="🗑 O'chirish", callback_data=f"rm:{r['user_id']}"),
        ])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


@router.message(F.text == BTN_ADMINS, F.chat.type == "private")
@router.message(Command("adminlar"), F.chat.type == "private")
async def cmd_admins(m: Message):
    if not is_owner(m.from_user.id):
        if is_admin(m.from_user.id):
            await m.answer("Bu bo'lim faqat bot egasi uchun.")
        else:
            await request_access(m)
        return
    rows = admin_rows()
    lines = ["👥 <b>Jamoa</b> (siz — 👑 egasi)\n"]
    for role in ("operator", "sotuvchi"):
        group = [r for r in rows if (r["role"] or "operator") == role]
        lines.append(f"<b>{ROLE_NAMES[role]}lar:</b>")
        lines += [f"  • {user_label(r['name'], r['username'])}" for r in group] or ["  — yo'q"]
    lines.append("\nYangi odam qo'shish: <b>➕ Operator qo'shish</b> yoki <b>➕ Sotuvchi qo'shish</b> ni bosing "
                 "va chiqqan havolani o'sha odamga yuboring.")
    await m.answer("\n".join(lines), reply_markup=team_kb(rows))


@router.callback_query(F.data.startswith("inv:"))
async def cb_invite(c: CallbackQuery):
    if not is_owner(c.from_user.id):
        await c.answer("Faqat bot egasi uchun", show_alert=True)
        return
    role = c.data[4:]
    token = secrets.token_urlsafe(8).replace("-", "x").replace("_", "y")
    execute("INSERT INTO invites(token, role, created_at) VALUES(?,?,?)", (token, role, now_local().strftime(FMT)))
    me = await c.bot.me()
    link = f"https://t.me/{me.username}?start=inv_{token}"
    await c.answer()
    # 1) egasi uchun yo'riqnoma
    await c.message.answer(
        f"✅ {ROLE_NAMES[role]} uchun taklif tayyor.\n\n"
        "👇 <b>Pastdagi xabarni</b> yangi xodimga <b>uzating (Forward)</b>. U tugmani bosib "
        f"<b>Start</b> ni bossa, darhol {ROLE_NAMES[role]} bo'lib qo'shiladi.\n\n"
        "⚠️ Taklif <b>bir martalik</b> va 24 soat amal qiladi.\n"
        f"Havolani nusxalash uchun: <code>{link}</code>")
    # 2) xodimga uzatiladigan xabar — tugma forward qilinganda ham ishlaydi
    kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="🚀 Botga qo'shilish", url=link)]])
    await c.message.answer(
        f"👋 Salom! Sizni <b>ANT</b> jamoasiga {ROLE_NAMES[role]} sifatida taklif qilishdi.\n\n"
        f"Qo'shilish uchun pastdagi <b>🚀 Botga qo'shilish</b> tugmasini bosing, keyin <b>Start</b> ni bosing.\n\n"
        f'Tugma ishlamasa, shu havolani bosing: <a href="{link}">ANT Tashrif boti</a>',
        reply_markup=kb)


@router.callback_query(F.data.startswith("role:"))
async def cb_switch_role(c: CallbackQuery):
    if not is_owner(c.from_user.id):
        await c.answer("Faqat bot egasi o'zgartira oladi", show_alert=True)
        return
    uid = int(c.data.split(":")[1])
    rows = q("SELECT * FROM admins WHERE user_id=?", (uid,))
    if not rows:
        await c.answer("Bu odam ro'yxatda yo'q", show_alert=True)
        return
    new_role = "sotuvchi" if (rows[0]["role"] or "operator") == "operator" else "operator"
    execute("UPDATE admins SET role=? WHERE user_id=?", (new_role, uid))
    await c.answer("O'zgartirildi")
    await c.message.answer(f"🔄 {user_label(rows[0]['name'], rows[0]['username'])} endi {ROLE_NAMES[new_role]}.")
    await send_private(c.bot, uid, f"Sizning rolingiz o'zgardi: <b>{ROLE_NAMES[new_role]}</b>\n\n" + help_for(uid),
                       kb_for(uid))


@router.callback_query(F.data.startswith("rm:"))
async def cb_remove_admin(c: CallbackQuery):
    if not is_owner(c.from_user.id):
        await c.answer("Faqat bot egasi o'chira oladi", show_alert=True)
        return
    uid = int(c.data.split(":")[1])
    rows = q("SELECT * FROM admins WHERE user_id=?", (uid,))
    if not rows:
        await c.answer("Allaqachon o'chirilgan", show_alert=True)
        return
    execute("DELETE FROM admins WHERE user_id=?", (uid,))
    execute("UPDATE requests SET status='olindi' WHERE user_id=?", (uid,))
    await c.answer("O'chirildi")
    await c.message.answer(f"🗑 {user_label(rows[0]['name'], rows[0]['username'])} jamoadan olindi.")
    await send_private(c.bot, uid, "Botdan foydalanish ruxsatingiz bekor qilindi.", ReplyKeyboardRemove())


# ================================================================ savdo summasi (sotuvchi yozadi)
def has_pending_amount(m: Message) -> bool:
    return bool(m.from_user) and bool(q("SELECT 1 FROM pending_amount WHERE user_id=?", (m.from_user.id,)))


@router.message(F.chat.type == "private", StateFilter(None), TEXT, has_pending_amount)
async def st_amount(m: Message):
    amount = parse_amount(m.text)
    if not amount:
        await m.answer("Summani tushunmadim. Masalan: <b>2500000</b> yoki <b>2,5 mln</b>")
        return
    vid = q("SELECT visit_id FROM pending_amount WHERE user_id=?", (m.from_user.id,))[0]["visit_id"]
    execute("UPDATE visits SET amount=? WHERE id=?", (amount, vid))
    execute("DELETE FROM pending_amount WHERE user_id=?", (m.from_user.id,))
    v = get_visit(vid)
    await m.answer(f"✅ Saqlandi: <b>{money(amount)}</b>", reply_markup=kb_for(m.from_user.id))
    await tell_creator(m.bot, v, f"💰 №{vid} — <b>{esc(v['client'])}</b>: savdo summasi {money(amount)} "
                                 f"({esc(m.from_user.full_name)})")


@router.message(F.chat.type == "private")
async def fallback(m: Message):
    if not is_admin(m.from_user.id):
        await request_access(m)
        return
    await m.answer(help_for(m.from_user.id), reply_markup=kb_for(m.from_user.id))


# ================================================================ eslatmalar va hisobotlar (fon)
def was_sent(vid, kind):
    return bool(q("SELECT 1 FROM sent WHERE visit_id=? AND kind=?", (vid, kind)))


def mark_sent(vid, kind):
    execute("INSERT OR IGNORE INTO sent(visit_id, kind) VALUES(?, ?)", (vid, kind))


def _hm(text):
    h, m = (int(x) for x in text.split(":"))
    return h, m


async def tick(bot, now=None):
    now = now or now_local()

    # 1) tashrifdan oldingi eslatmalar
    for v in q("SELECT * FROM visits WHERE status='kutilmoqda' AND visit_at > ?", (now.strftime(FMT),)):
        vdt = datetime.strptime(v["visit_at"], FMT)
        created = datetime.strptime(v["created_at"], FMT)
        due = [mn for mn in REMINDERS
               if now >= vdt - timedelta(minutes=mn) and not was_sent(v["id"], f"r{mn}")]
        if not due:
            continue
        if any(vdt - timedelta(minutes=mn) >= created for mn in due):
            left = (vdt - now).total_seconds() // 60
            if not await notify_visit(bot, v["id"], f"⏰ <b>Eslatma: mijoz kelishiga {human_minutes(left)} qoldi</b>"):
                continue
        for mn in due:
            mark_sent(v["id"], f"r{mn}")

    # 2) "mijoz keldimi?" — vaqt o'tdi, hali belgilanmagan
    cutoff = (now - timedelta(minutes=CHECK_AFTER_MINUTES)).strftime(FMT)
    for v in q("SELECT * FROM visits WHERE status='kutilmoqda' AND visit_at <= ?", (cutoff,)):
        if was_sent(v["id"], "chk"):
            continue
        if now - datetime.strptime(v["visit_at"], FMT) > timedelta(days=1):
            mark_sent(v["id"], "chk")  # juda eski — eslatmaymiz
            continue
        if await notify_visit(bot, v["id"], "❓ <b>Mijoz keldimi?</b> Tashrif vaqti o'tdi — belgilang 👇"):
            mark_sent(v["id"], "chk")

    today = now.date().isoformat()

    # 3) ertalabki ro'yxat — butun jamoaga shaxsan
    if (now.hour, now.minute) >= _hm(DAILY_SUMMARY_TIME) and get_meta("summary") != today:
        set_meta("summary", today)
        text = day_summary(now.date())
        if text:
            for uid in set(team_ids("sotuvchi", "operator") + ([owner_id()] if owner_id() else [])):
                await send_private(bot, uid, "☀️ Xayrli tong!\n\n" + text)

    # 4) kechki hisobot (egasiga)
    if (now.hour, now.minute) >= _hm(REPORT_TIME) and get_meta("evening") != today:
        owner = owner_id()
        if owner:
            await send_private(bot, owner, "🌙 <b>Kunlik hisobot</b>\n\n" + build_report(now.date(), now.date()),
                               report_kb())
        set_meta("evening", today)


async def scheduler(bot):
    while True:
        try:
            await tick(bot)
        except Exception:
            logging.exception("Fon tekshiruvida xato")
        await asyncio.sleep(30)


# ================================================================ ishga tushirish
async def main():
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    if not BOT_TOKEN:
        raise SystemExit(".env faylida BOT_TOKEN yozilmagan")
    init_db()
    if owner_id():
        logging.info("Bot egasi: %s", owner_id())
    else:
        logging.warning("Bot egasi hali yo'q: botga birinchi /start yozgan odam egasi bo'ladi")
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        logging.warning("openpyxl o'rnatilmagan — Excel o'rniga CSV beriladi. O'rnatish: pip install -r requirements.txt")
    bot = Bot(BOT_TOKEN, default=DefaultBotProperties(parse_mode="HTML"))
    dp = Dispatcher(storage=MemoryStorage())
    dp.include_router(router)
    await bot.set_my_commands([
        BotCommand(command="start", description="Menyuni ochish"),
        BotCommand(command="yangi", description="Yangi tashrif (operator)"),
        BotCommand(command="bugun", description="Bugungi tashriflar"),
        BotCommand(command="royxat", description="Kelgusi tashriflar"),
        BotCommand(command="hisobot", description="Hisobot (egasi)"),
        BotCommand(command="adminlar", description="Jamoa (egasi)"),
        BotCommand(command="bekor", description="Kiritishni to'xtatish"),
        BotCommand(command="id", description="Chat ID"),
    ])
    reminder_task = asyncio.create_task(scheduler(bot))
    try:
        await dp.start_polling(bot)
    finally:
        reminder_task.cancel()


if __name__ == "__main__":
    asyncio.run(main())
